"""Shared swim-lane slide-building helpers, pulled out of
build_unpaid_leave_slide.py once a second slide needed the exact same
primitives - colors/typography read directly from the real deck (slide 3),
not eyeballed. Every new feature/system slide should build on these so the
whole deck stays visually consistent without copy-pasting ~300 lines each."""

from lxml import etree

from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

NAVY = RGBColor(0x1F, 0x2A, 0x4A)
EMPLOYEE_FILL, EMPLOYEE_LINE = RGBColor(0xE6, 0xEE, 0xF9), RGBColor(0x93, 0xAE, 0xD6)
AGENT_FILL, AGENT_LINE = RGBColor(0xF1, 0xEA, 0xFC), RGBColor(0x9A, 0x74, 0xDC)
TOOL_FILL, TOOL_LINE = RGBColor(0xF2, 0xF4, 0xF7), RGBColor(0xA7, 0xB2, 0xBF)
HR_FILL, HR_LINE = RGBColor(0xBE, 0x18, 0x5D), RGBColor(0x9D, 0x17, 0x4D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE_DARK, SLATE_MED = RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x47, 0x55, 0x69)
PINK_TEXT = RGBColor(0xFC, 0xE7, 0xF3)
ARROW_LINE = RGBColor(0x33, 0x4E, 0x7A)
PANEL_BG = RGBColor(0xF6, 0xF8, 0xFA)
PANEL_BORDER = RGBColor(0xDB, 0x7A, 0xA5)
NOTE_MAGENTA = RGBColor(0x9D, 0x17, 0x4D)
STORE_FILL, STORE_LINE = RGBColor(0xDC, 0xEF, 0xE3), RGBColor(0x8B, 0xC3, 0xA0)

FONT = "Calibri"


def box(slide, x, y, w, h, fill, line, title, body, title_color=SLATE_DARK, body_color=SLATE_MED, title_size=8.5, body_size=6.8):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(45720)
    tf.margin_right = Emu(45720)
    tf.margin_top = Emu(27432)
    tf.margin_bottom = Emu(27432)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = title
    r1.font.bold = True
    r1.font.size = Pt(title_size)
    r1.font.name = FONT
    r1.font.color.rgb = title_color
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = body
    r2.font.size = Pt(body_size)
    r2.font.name = FONT
    r2.font.color.rgb = body_color
    return shape


def arrow(slide, x1, y1, x2, y2):
    """A plain shape whose prstGeom is patched to "line" with a <a:tailEnd
    triangle> - matches exactly how the real deck draws every arrow
    (confirmed from its own XML). python-pptx's MSO_SHAPE enum has no bare
    LINE member, so build a RECTANGLE and rewrite its prstGeom in place -
    a Connector object doesn't expose arrowheads in python-pptx at all."""
    width = abs(x2 - x1) or 1
    height = abs(y2 - y1) or 1
    line_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(min(x1, x2)), Emu(min(y1, y2)), Emu(width), Emu(height))
    prst_geom = line_shape._element.spPr.find(qn("a:prstGeom"))
    prst_geom.set("prst", "line")
    for av_lst in prst_geom.findall(qn("a:avLst")):
        prst_geom.remove(av_lst)
    line_shape.fill.background()
    line_shape.line.color.rgb = ARROW_LINE
    line_shape.line.width = Pt(1.4)
    line_shape.shadow.inherit = False
    ln = line_shape.line._get_or_add_ln()
    tail_end = etree.SubElement(ln, f"{{{A_NS}}}tailEnd")
    tail_end.set("type", "triangle")
    return line_shape


def legend_chip(slide, x, y, w, fill, line, label):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(114300))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.5)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = Emu(18288)
    tf.margin_right = Emu(18288)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(7.5)
    r.font.name = FONT
    r.font.color.rgb = SLATE_DARK


def new_slide(prs):
    """Blank slide, same layout slide 18 uses - strips any placeholder
    shapes the layout would otherwise add."""
    blank_layout = prs.slides[2].slide_layout
    slide = prs.slides.add_slide(blank_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    return slide


def add_header(slide, title_text, subtitle_text):
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), Emu(12191695), Emu(482600))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.color.rgb = NAVY
    header.shadow.inherit = False
    tf = header.text_frame
    tf.margin_left = Emu(177800)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title_text
    r.font.bold = True
    r.font.size = Pt(19)
    r.font.name = FONT
    r.font.color.rgb = WHITE

    subtitle = slide.shapes.add_textbox(Emu(6500000), Emu(50800), Emu(5600000), Emu(381000))
    stf = subtitle.text_frame
    stf.word_wrap = True
    stf.vertical_anchor = MSO_ANCHOR.MIDDLE
    sp = stf.paragraphs[0]
    sp.alignment = PP_ALIGN.RIGHT
    sr = sp.add_run()
    sr.text = subtitle_text
    sr.font.size = Pt(9)
    sr.font.name = FONT
    sr.font.color.rgb = WHITE


def add_section_label(slide, text, y=635000, width=7400000):
    label = slide.shapes.add_textbox(Emu(228600), Emu(y), Emu(width), Emu(165100))
    lp = label.text_frame.paragraphs[0]
    lr = lp.add_run()
    lr.text = text
    lr.font.bold = True
    lr.font.size = Pt(10.5)
    lr.font.name = FONT
    lr.font.color.rgb = SLATE_DARK


def even_columns(start_x, total_width, n, gap=70000):
    """n evenly-spaced column x-positions plus their shared width, so a new
    slide's box row doesn't need hand-computed EMU math like slide 18's did."""
    box_w = (total_width - (n - 1) * gap) // n
    xs = [start_x + i * (box_w + gap) for i in range(n)]
    return xs, box_w


def swim_lane(slide, label_text, panel_y, top_row, bottom_row, panel_h=1700000):
    """top_row: list of (x, w, fill, line, title, body, tcolor, bcolor).
    bottom_row: list of (x, w, title, body) - always STORE_FILL styled.
    Draws the outer band panel, the label, both rows, and one down arrow
    per top box into whichever bottom box sits under it."""
    band_panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(127000), Emu(panel_y), Emu(11938000), Emu(panel_h))
    band_panel.fill.solid()
    band_panel.fill.fore_color.rgb = WHITE
    band_panel.line.color.rgb = PANEL_BORDER
    band_panel.line.width = Pt(1.25)
    band_panel.shadow.inherit = False
    band_panel.text_frame.text = ""

    label = slide.shapes.add_textbox(Emu(228600), Emu(panel_y + 30000), Emu(9000000), Emu(165100))
    lp = label.text_frame.paragraphs[0]
    lr = lp.add_run()
    lr.text = label_text
    lr.font.bold = True
    lr.font.size = Pt(11)
    lr.font.name = FONT
    lr.font.color.rgb = SLATE_DARK

    top_y = panel_y + 230000
    top_h = 460000
    for bx, bw_, fill, line, title, body, tcolor, bcolor in top_row:
        box(slide, bx, top_y, bw_, top_h, fill, line, title, body, title_color=tcolor, body_color=bcolor, title_size=8.5, body_size=6.8)

    bottom_y = top_y + top_h + 100000
    bottom_h = 420000
    for bx, bw_, title, body in bottom_row:
        box(slide, bx, bottom_y, bw_, bottom_h, STORE_FILL, STORE_LINE, title, body, title_color=SLATE_DARK, body_color=SLATE_MED, title_size=8, body_size=6.5)

    for bx, bw_, *_rest in top_row:
        center_x = bx + bw_ // 2
        for store_x, store_w, *_ in bottom_row:
            if store_x <= center_x <= store_x + store_w:
                arrow(slide, center_x, top_y + top_h, center_x, bottom_y)
                break

    return panel_y + panel_h


def bullet_panel(slide, x, y, w, h, fill, line, heading, bullets, heading_color=SLATE_DARK, bullet_color=SLATE_MED):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = fill
    panel.line.color.rgb = line
    panel.line.width = Pt(1)
    panel.shadow.inherit = False
    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(91440)
    tf.margin_top = Emu(45720)
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = heading
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.name = FONT
    r1.font.color.rgb = heading_color
    for b in bullets:
        bp = tf.add_paragraph()
        br = bp.add_run()
        br.text = "• " + b
        br.font.size = Pt(8.5)
        br.font.name = FONT
        br.font.color.rgb = bullet_color
    return panel


def save_deck(prs, dst, dst_tmp):
    try:
        prs.save(dst)
        print("Saved:", dst)
    except PermissionError:
        prs.save(dst_tmp)
        print("dst is open/locked - saved work-in-progress copy instead:", dst_tmp)
